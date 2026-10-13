"""
Document loaders module for the RAG Chatbot Pipeline.
Supports loading PDF, DOCX, TXT, and Markdown files from disk or uploaded file streams.
"""

from __future__ import annotations
import os
import io
from pathlib import Path
from typing import List, Union, BinaryIO
import logging

from langchain_core.documents import Document

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv"}


def load_txt_or_md(file_path: Union[str, Path]) -> List[Document]:
    """Load plain text or markdown file."""
    path = Path(file_path)
    content = ""
    for enc in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
        try:
            with open(path, "r", encoding=enc) as f:
                content = f.read()
            break
        except UnicodeDecodeError:
            continue

    metadata = {
        "source": str(path.name),
        "file_path": str(path.resolve()),
        "file_type": path.suffix.lower(),
        "total_chars": len(content),
    }
    return [Document(page_content=content, metadata=metadata)]


def load_pdf(file_path: Union[str, Path]) -> List[Document]:
    """Load PDF file and extract pages using pypdf."""
    path = Path(file_path)
    import pypdf

    docs: List[Document] = []
    reader = pypdf.PdfReader(str(path))
    total_pages = len(reader.pages)

    for idx, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            metadata = {
                "source": str(path.name),
                "file_path": str(path.resolve()),
                "file_type": ".pdf",
                "page": idx + 1,
                "total_pages": total_pages,
            }
            docs.append(Document(page_content=text, metadata=metadata))

    if not docs:
        logger.warning("No extractable text found in PDF: %s", path.name)
    return docs


def load_docx(file_path: Union[str, Path]) -> List[Document]:
    """Load DOCX file using python-docx."""
    path = Path(file_path)
    import docx

    doc = docx.Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    content = "\n\n".join(paragraphs)

    metadata = {
        "source": str(path.name),
        "file_path": str(path.resolve()),
        "file_type": ".docx",
        "paragraphs_count": len(paragraphs),
    }
    return [Document(page_content=content, metadata=metadata)]


def load_document(file_path: Union[str, Path]) -> List[Document]:
    """Load a single document by its file extension."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = path.suffix.lower()
    if ext == ".pdf":
        return load_pdf(path)
    elif ext == ".docx":
        return load_docx(path)
    elif ext in {".txt", ".md", ".csv"}:
        return load_txt_or_md(path)
    else:
        raise ValueError(
            f"Unsupported file format '{ext}'. Supported formats: {SUPPORTED_EXTENSIONS}"
        )


def load_directory(
    dir_path: Union[str, Path], recursive: bool = True
) -> List[Document]:
    """Load all supported documents from a directory."""
    path = Path(dir_path)
    if not path.is_dir():
        raise NotADirectoryError(f"Directory not found: {dir_path}")

    documents: List[Document] = []
    pattern = "**/*" if recursive else "*"

    for item in path.glob(pattern):
        if item.is_file() and item.suffix.lower() in SUPPORTED_EXTENSIONS:
            try:
                docs = load_document(item)
                documents.extend(docs)
                logger.info("Loaded '%s' (%d doc segments)", item.name, len(docs))
            except Exception as exc:
                logger.error("Failed to load '%s': %s", item.name, exc)

    return documents


def load_from_uploaded_file(
    file_name: str, file_bytes: bytes
) -> List[Document]:
    """Extract documents directly from uploaded file bytes (for Streamlit/API)."""
    ext = Path(file_name).suffix.lower()

    if ext == ".pdf":
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        total_pages = len(reader.pages)
        docs = []
        for idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                docs.append(
                    Document(
                        page_content=text,
                        metadata={
                            "source": file_name,
                            "file_type": ".pdf",
                            "page": idx + 1,
                            "total_pages": total_pages,
                        },
                    )
                )
        return docs

    elif ext == ".docx":
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        content = "\n\n".join(paragraphs)
        return [
            Document(
                page_content=content,
                metadata={"source": file_name, "file_type": ".docx"},
            )
        ]

    elif ext in {".txt", ".md", ".csv"}:
        text = ""
        for enc in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                text = file_bytes.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        return [
            Document(
                page_content=text,
                metadata={"source": file_name, "file_type": ext},
            )
        ]
    else:
        raise ValueError(f"Unsupported uploaded file format: {ext}")
